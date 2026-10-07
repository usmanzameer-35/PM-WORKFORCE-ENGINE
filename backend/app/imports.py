import csv, io, json, zipfile
from openpyxl import load_workbook
from .schemas import Portfolio, State

TEMPLATE = "id,name,region,property_type,doors,assigned_to,properties,clients,monthly_revenue,maintenance,new_leases,inspections\nimport-1,Example Portfolio,Toronto,Multifamily,100,pm-6,2,1,10500,15,2,6\n"


def preview(
    data: bytes, filename: str, state: State, mapping: dict[str, str] | None = None
):
    if len(data) > 2_000_000:
        raise ValueError("Upload limit is 2 MB")
    if filename.lower().endswith(".xlsx"):
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            if sum(f.file_size for f in archive.infolist()) > 20_000_000:
                raise ValueError("Expanded workbook exceeds 20 MB")
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        sheet = wb.active
        iterator = sheet.iter_rows(values_only=True)
        headers = next(iterator, None)
        if not headers:
            raise ValueError("Workbook is empty")
        headers = [str(h).strip() for h in headers]
        rows = []
        for row in iterator:
            if any(v is not None for v in row):
                rows.append(dict(zip(headers, row)))
            if len(rows) > 1000:
                raise ValueError("Maximum 1,000 rows per import")
        wb.close()
    elif filename.lower().endswith(".csv"):
        reader = csv.DictReader(io.StringIO(data.decode("utf-8-sig")))
        headers = reader.fieldnames or []
        rows = list(reader)
    else:
        raise ValueError("Use a CSV or XLSX file")
    if len(rows) > 1000:
        raise ValueError("Maximum 1,000 rows per import")
    if len(set(headers)) != len(headers):
        raise ValueError("Duplicate column names")
    mapping = mapping or {h: h for h in headers}
    allowed = set(Portfolio.model_fields) | set(state.settings.coefficients) - {"doors"}
    if any(v not in allowed for v in mapping.values()):
        raise ValueError("Mapping contains an unknown destination")
    if len(set(mapping.values())) != len(mapping):
        raise ValueError("Multiple columns map to the same destination")
    existing = {p.id for p in state.portfolios}
    seen = set()
    accepted = []
    errors = []
    for number, row in enumerate(rows, 2):
        try:
            row = {
                mapping[k]: v
                for k, v in row.items()
                if k in mapping and v is not None and str(v).strip() != ""
            }
            activities = {
                k: row.pop(k)
                for k in list(row)
                if k in state.settings.coefficients and k != "doors"
            }
            if "activities" in row:
                activities.update(json.loads(row.pop("activities")))
            if "complexity" in row:
                row["complexity"] = json.loads(row["complexity"])
            p = Portfolio(**row, activities=activities)
            if p.id in existing or p.id in seen:
                raise ValueError("ID already exists; imports only add new portfolios")
            State(**{**state.model_dump(), "portfolios": [p.model_dump()]})
            accepted.append(p.model_dump())
            seen.add(p.id)
        except (ValueError, TypeError) as e:
            errors.append({"row": number, "message": str(e)[:600]})
    if not rows:
        errors.append({"row": 1, "message": "No data rows found"})
    return {
        "headers": headers,
        "rows": accepted,
        "errors": errors,
        "count": len(rows),
        "version": state.version,
        "can_commit": bool(rows) and not errors,
        "note": "Missing activity and complexity inputs default to zero. Review demand assumptions before committing.",
    }
